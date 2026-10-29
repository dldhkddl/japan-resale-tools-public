# -*- coding: utf-8 -*-
"""라쿠텐 일본 수요 자동 수집기
설정.json의 키워드로 라쿠텐 상품을 리뷰 많은 순으로 모아
역계산 칸이 붙은 엑셀(일본수요_날짜.xlsx)을 만든다.
"""
import json, os, re, sys, time, datetime, urllib.parse, urllib.request, urllib.error

BASE = os.path.dirname(os.path.abspath(__file__))
API = "https://openapi.rakuten.co.jp/ichibams/api/IchibaItem/Search/20260701"


def load_config():
    with open(os.path.join(BASE, "설정.json"), encoding="utf-8") as f:
        cfg = json.load(f)
    if "여기에" in cfg["applicationId"] or "여기에" in cfg["accessKey"]:
        sys.exit("설정.json에 applicationId와 accessKey를 먼저 넣어주세요.")
    return cfg


def call_api(cfg, keyword, page):
    params = {
        "applicationId": cfg["applicationId"],
        "accessKey": cfg["accessKey"],
        "keyword": keyword,
        "sort": "-reviewCount",
        "hits": 30,
        "page": page,
        "formatVersion": 2,
        "availability": 1,
        "hasReviewFlag": 1,
    }
    url = API + "?" + urllib.parse.urlencode(params)
    headers = {"User-Agent": "japan-demand-collector/1.0", "accessKey": cfg["accessKey"]}
    if cfg.get("referer"):
        headers["Referer"] = cfg["referer"]
        headers["Origin"] = cfg["referer"].rstrip("/")
    for attempt in range(3):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=20) as r:
                return json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            body = e.read().decode("utf-8", "ignore")[:300]
            if e.code == 429:
                print("  요청이 많아 10초 쉬었다 다시 시도…")
                time.sleep(10)
                continue
            if e.code == 403 and "REFER" in body.upper():
                sys.exit("403 Referer 오류: 라쿠텐 앱에 등록한 사이트 주소를 설정.json의 referer에 넣어주세요.\n" + body)
            if e.code == 404 and page > 1:
                return None  # 결과 페이지 끝
            sys.exit(f"API 오류 {e.code}: {body}")
        except urllib.error.URLError as e:
            print(f"  연결 오류({e.reason}), 다시 시도…")
            time.sleep(3)
    sys.exit("API 호출에 계속 실패했어요. 잠시 뒤 다시 실행해 주세요.")


def extract_items(data):
    if not data:
        return []
    raw = data.get("items") or data.get("Items") or []
    out = []
    for it in raw:
        if isinstance(it, dict) and "item" in it:  # formatVersion 1 대비
            it = it["item"]
        if isinstance(it, dict) and "Item" in it:
            it = it["Item"]
        out.append(it)
    return out


def image_of(it):
    imgs = it.get("mediumImageUrls") or []
    if imgs:
        first = imgs[0]
        return first.get("imageUrl", "") if isinstance(first, dict) else str(first)
    return ""


BRACKETS = re.compile(r"【([^】]*)】|\[([^\]]*)\]|［([^］]*)］|≪([^≫]*)≫")
PROMO = re.compile(r"OFF|ＯＦＦ|クーポン|％|%|円|送料|ポイント|SALE|セール|倍|まで|限定|楽天|ランキング|1位|１位|レビュー|即納|予約")


def clean_name(name, limit=90):
    """판촉 문구 괄호는 지우고, 상품명이 든 괄호는 괄호만 벗긴다."""
    def repl(m):
        inner = next(g for g in m.groups() if g is not None)
        return " " if PROMO.search(inner) else " " + inner + " "
    t = BRACKETS.sub(repl, name)
    t = re.sub(r"★[^★]*★", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t[:limit]


def load_cache():
    try:
        with open(os.path.join(BASE, "번역캐시.json"), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_cache(cache):
    with open(os.path.join(BASE, "번역캐시.json"), "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=0)


def translate_google_free(text):
    q = urllib.parse.urlencode({"client": "gtx", "sl": "ja", "tl": "ko", "dt": "t", "q": text})
    req = urllib.request.Request("https://translate.googleapis.com/translate_a/single?" + q,
                                 headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=15) as r:
        data = json.loads(r.read().decode("utf-8"))
    return "".join(seg[0] for seg in data[0] if seg and seg[0])


def translate_deepl(text, key):
    host = "api-free.deepl.com" if key.endswith(":fx") else "api.deepl.com"
    body = urllib.parse.urlencode({"text": text, "source_lang": "JA", "target_lang": "KO"}).encode()
    req = urllib.request.Request(f"https://{host}/v2/translate", data=body,
                                 headers={"Authorization": "DeepL-Auth-Key " + key})
    with urllib.request.urlopen(req, timeout=15) as r:
        return json.loads(r.read().decode("utf-8"))["translations"][0]["text"]


def add_korean_names(cfg, rows):
    mode = cfg.get("translate", "google_free")
    if mode == "off":
        return
    cache = load_cache()
    todo = [v for v in rows if clean_name(v["name"]) not in cache]
    if todo:
        print(f"\n한국어 상품명 번역 중… ({len(todo)}개, 이전에 번역한 건 건너뜀)")
    fails = 0
    for i, v in enumerate(todo, 1):
        src = clean_name(v["name"])
        try:
            if mode == "deepl":
                cache[src] = translate_deepl(src, cfg.get("deeplKey", ""))
            else:
                cache[src] = translate_google_free(src)
            time.sleep(0.3)
        except Exception as e:
            fails += 1
            if fails == 1:
                print(f"  번역 실패({e}). 실패한 칸은 비워두고 계속해요.")
            if fails >= 5 and fails == i:
                print("  번역이 계속 막혀서 건너뜀. 설정.json의 translate를 \"deepl\"로 바꾸고 deeplKey를 넣어보세요.")
                break
        if i % 20 == 0:
            print(f"  {i}/{len(todo)}")
            save_cache(cache)
    save_cache(cache)
    for v in rows:
        v["name_ko"] = cache.get(clean_name(v["name"]), "")


GENERIC_KO = {"인기", "랭킹", "추천", "무료배송", "송료무료", "선물", "선택", "선택가능", "정규품", "공식", "한국", "한국화장품",
              "코스메", "개별포장", "대용량", "세트", "신상품", "당일발송", "즉납", "점장", "메일", "우편", "배송"}


def make_query(v, words=6):
    base = v.get("name_ko") or clean_name(v["name"])
    base = re.sub(r"[「」『』()（）\[\]【】★☆!！?？,，、・/|]", " ", base)
    toks = [t for t in base.split() if t not in GENERIC_KO and not re.fullmatch(r"\d+(위|%|엔|円)?", t)]
    return " ".join(toks[:words])


def search_urls(q):
    e = urllib.parse.quote(q)
    return {"naver": "https://search.shopping.naver.com/search/all?query=" + e,
            "danawa": "https://search.danawa.com/dsearch.php?query=" + e,
            "coupang": "https://www.coupang.com/np/search?q=" + e}


def collect(cfg):
    items = {}
    for kw in cfg["keywords"]:
        print(f"[{kw}] 수집 중…")
        for page in range(1, int(cfg.get("pagesPerKeyword", 2)) + 1):
            data = call_api(cfg, kw, page)
            got = extract_items(data)
            if not got:
                break
            for it in got:
                code = it.get("itemCode")
                if not code:
                    continue
                if code in items:
                    if kw not in items[code]["keywords"]:
                        items[code]["keywords"].append(kw)
                    continue
                items[code] = {
                    "keywords": [kw],
                    "name": it.get("itemName", ""),
                    "price": it.get("itemPrice", 0),
                    "reviews": it.get("reviewCount", 0) or 0,
                    "rating": it.get("reviewAverage", 0) or 0,
                    "shop": it.get("shopName", ""),
                    "postage": "O" if it.get("postageFlag") == 0 else "X",
                    "url": it.get("itemUrl", ""),
                    "image": image_of(it),
                }
            print(f"  {page}페이지: {len(got)}개")
            time.sleep(1.2)  # 과호출 방지
            if data.get("pageCount") and page >= data["pageCount"]:
                break
    rows = [v for v in items.values() if float(v["reviews"]) >= cfg.get("minReviewCount", 0)]
    rows.sort(key=lambda v: float(v["reviews"]), reverse=True)
    return rows


def write_xlsx(cfg, rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation

    FN = "맑은 고딕"
    f = lambda **k: Font(name=FN, **k)
    YEL = PatternFill("solid", fgColor="FFF2CC")
    HEAD = PatternFill("solid", fgColor="1F3864")
    th = Side(style="thin", color="BFBFBF")
    BD = Border(left=th, right=th, top=th, bottom=th)
    wb = Workbook()

    # 설정
    s = wb.active
    s.title = "설정"
    s["A1"] = "역계산 설정 (노란 칸 수정 → 수집결과 시트 자동 재계산)"
    s["A1"].font = f(bold=True, size=13)
    conf = [("환율 (원/엔)", cfg["exchangeRate_KRW_per_JPY"], "0.00", "오늘 환율로 바꾸세요"),
            ("플랫폼 수수료율", cfg["platformFeeRate"], "0.0%", "입점할 플랫폼 수수료 확인 후 입력"),
            ("환율 여유분", cfg["fxBufferRate"], "0.0%", "환율 변동 대비"),
            ("목표 마진율", cfg["targetMarginRate"], "0.0%", "이 이상이면 판매"),
            ("상품당 포장비 (원)", cfg["packingCost_KRW"], "#,##0", "예시값")]
    for i, (k, v, fm, note) in enumerate(conf, start=3):
        s.cell(i, 1, k).font = f(bold=True)
        c = s.cell(i, 2, v); c.font = f(color="0000FF"); c.fill = YEL; c.number_format = fm; c.border = BD
        s.cell(i, 3, note).font = f(color="7F7F7F")
    s["A10"] = "국제배송비 표 (무게 하한 g → 배송비 원, 예시값)"; s["A10"].font = f(bold=True)
    ship = [(0, 5500), (500, 7500), (1000, 10500), (2000, 15000), (3000, 19500)]
    for i, (lo, cost) in enumerate(ship, start=11):
        for col, v in ((1, lo), (2, cost)):
            c = s.cell(i, col, v); c.font = f(color="0000FF"); c.fill = YEL; c.border = BD; c.number_format = "#,##0"
    s["A17"] = f"수집일: {datetime.date.today()} / 키워드: {', '.join(cfg['keywords'])}"
    s["A17"].font = f(color="7F7F7F")
    s.column_dimensions["A"].width = 22; s.column_dimensions["B"].width = 12; s.column_dimensions["C"].width = 40

    # 수집결과
    w = wb.create_sheet("수집결과", 0)
    w["A1"] = "라쿠텐 일본 수요 (리뷰 많은 순)"; w["A1"].font = f(bold=True, size=13)
    w["A2"] = "노란 칸(한국 구매가·무게·성분 확인)만 채우면 판정이 나와요. 판매가는 라쿠텐 표시가 그대로라 샵 배송비는 별도일 수 있어요."
    w["A2"].font = f(italic=True, color="7F7F7F")
    from openpyxl.utils import get_column_letter as L
    cols = ["순위", "키워드", "상품명 (일본어)", "한국어 상품명", "일본 판매가 (엔)", "리뷰 수", "평점", "샵", "배송비 포함", "라쿠텐",
            "한국 검색어 (수정 가능)", "네이버쇼핑", "다나와", "쿠팡",
            "한국 구매가 (원)", "무게 (g)", "일본 성분 확인",
            "판매가 (원)", "수수료 (원)", "배송비 (원)", "포장+환율여유 (원)", "순수익 (원)", "마진율", "판정", "최대 구매가 (원)"]
    C = {name: i + 1 for i, name in enumerate(cols)}
    X = {k: L(v) for k, v in C.items()}
    for j, h in enumerate(cols, 1):
        c = w.cell(4, j, h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD; c.border = BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    w.row_dimensions[4].height = 32
    LINK = f(color="0563C1", underline="single")
    for i, v in enumerate(rows):
        r = 5 + i
        vals = [i + 1, ", ".join(v["keywords"]), v["name"], v.get("name_ko", ""), v["price"], v["reviews"], v["rating"], v["shop"], v["postage"]]
        for j, val in enumerate(vals, 1):
            c = w.cell(r, j, val); c.font = f(); c.border = BD
        lc = w.cell(r, C["라쿠텐"], "열기"); lc.hyperlink = v["url"]; lc.font = LINK; lc.border = BD
        q = w.cell(r, C["한국 검색어 (수정 가능)"], make_query(v)); q.fill = YEL; q.font = f(color="0000FF"); q.border = BD
        Q = f'{X["한국 검색어 (수정 가능)"]}{r}'
        for name, base, label in (("네이버쇼핑", "https://search.shopping.naver.com/search/all?query=", "네이버"),
                                  ("다나와", "https://search.danawa.com/dsearch.php?query=", "다나와"),
                                  ("쿠팡", "https://www.coupang.com/np/search?q=", "쿠팡")):
            c = w.cell(r, C[name], f'=IF({Q}="","",HYPERLINK("{base}"&_xlfn.ENCODEURL({Q}),"{label}"))')
            c.font = LINK; c.border = BD; c.alignment = Alignment(horizontal="center")
        for name in ("한국 구매가 (원)", "무게 (g)", "일본 성분 확인"):
            c = w.cell(r, C[name]); c.fill = YEL; c.font = f(color="0000FF"); c.border = BD
        buy, wt, ok = (f'{X["한국 구매가 (원)"]}{r}', f'{X["무게 (g)"]}{r}', f'{X["일본 성분 확인"]}{r}')
        jp, sell = f'{X["일본 판매가 (엔)"]}{r}', f'{X["판매가 (원)"]}{r}'
        fee, ship, pack = f'{X["수수료 (원)"]}{r}', f'{X["배송비 (원)"]}{r}', f'{X["포장+환율여유 (원)"]}{r}'
        prof, mg = f'{X["순수익 (원)"]}{r}', f'{X["마진율"]}{r}'
        fm = {"판매가 (원)": f'={jp}*설정!$B$3',
              "수수료 (원)": f'={sell}*설정!$B$4',
              "배송비 (원)": f'=IF({wt}="","",INDEX(설정!$B$11:$B$15,MATCH({wt},설정!$A$11:$A$15,1)))',
              "포장+환율여유 (원)": f'=설정!$B$7+{sell}*설정!$B$5',
              "순수익 (원)": f'=IF(OR({buy}="",{wt}=""),"",{sell}-{buy}-{fee}-{ship}-{pack})',
              "마진율": f'=IF({prof}="","",{prof}/{sell})',
              "판정": f'=IF({prof}="","입력 필요",IF({ok}<>"O","보류",IF({mg}>=설정!$B$6,"판매","패스")))',
              "최대 구매가 (원)": f'=IF({wt}="","",{sell}*(1-설정!$B$4-설정!$B$5-설정!$B$6)-{ship}-설정!$B$7)'}
        for name, val in fm.items():
            c = w.cell(r, C[name], val); c.font = f(); c.border = BD
        for name in ("일본 판매가 (엔)", "리뷰 수", "한국 구매가 (원)", "무게 (g)", "판매가 (원)", "수수료 (원)", "배송비 (원)",
                     "포장+환율여유 (원)", "순수익 (원)", "최대 구매가 (원)"):
            w.cell(r, C[name]).number_format = "#,##0;(#,##0);-"
        w.cell(r, C["평점"]).number_format = "0.00"; w.cell(r, C["마진율"]).number_format = "0.0%"
        for name in ("순위", "배송비 포함", "일본 성분 확인", "판정"):
            w.cell(r, C[name]).alignment = Alignment(horizontal="center")
        for name in ("상품명 (일본어)", "한국어 상품명"):
            w.cell(r, C[name]).alignment = Alignment(wrap_text=True, vertical="top")
    last = 4 + max(len(rows), 1)
    dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True); w.add_data_validation(dv)
    dv.add(f'{X["일본 성분 확인"]}5:{X["일본 성분 확인"]}{last}')
    rng = f'{X["판정"]}5:{X["판정"]}{last}'
    w.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"판매"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FN, color="006100", bold=True)))
    w.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"패스"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=FN, color="9C0006")))
    w.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"보류"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=FN, color="9C5700")))
    widths = {"순위": 6, "키워드": 14, "상품명 (일본어)": 36, "한국어 상품명": 36, "샵": 16, "한국 검색어 (수정 가능)": 28,
              "네이버쇼핑": 8, "다나와": 8, "쿠팡": 8, "라쿠텐": 7}
    for name, col in C.items():
        w.column_dimensions[L(col)].width = widths.get(name, 11)
    w.freeze_panes = "E5"; w.auto_filter.ref = f"A4:{L(len(cols))}{last}"
    wb.save(path)


COLS = ["순위", "키워드", "상품명 (일본어)", "한국어 상품명", "일본 판매가 (엔)", "리뷰 수", "평점", "샵", "배송비 포함", "라쿠텐",
        "한국 검색어 (수정 가능)", "네이버쇼핑", "다나와", "쿠팡",
        "한국 구매가 (원)", "무게 (g)", "일본 성분 확인",
        "판매가 (원)", "수수료 (원)", "배송비 (원)", "포장+환율여유 (원)", "순수익 (원)", "마진율", "판정", "최대 구매가 (원)"]
SEARCH = (("네이버쇼핑", "https://search.shopping.naver.com/search/all?query=", "네이버"),
          ("다나와", "https://search.danawa.com/dsearch.php?query=", "다나와"),
          ("쿠팡", "https://www.coupang.com/np/search?q=", "쿠팡"))


def write_xlsx_lite(cfg, rows, path):
    """openpyxl이 없을 때 쓰는 표준 라이브러리 버전 (링크·수식 동작, 조건부 색상은 없음)."""
    import xlsx_lite as XL
    C = {n: i + 1 for i, n in enumerate(COLS)}
    X = {n: XL.col_letter(i) for n, i in C.items()}
    conf = [("환율 (원/엔)", cfg["exchangeRate_KRW_per_JPY"], "오늘 환율로 바꾸세요"),
            ("플랫폼 수수료율", cfg["platformFeeRate"], "입점할 플랫폼 수수료 확인 후 입력 (0.12 = 12%)"),
            ("환율 여유분", cfg["fxBufferRate"], "환율 변동 대비"),
            ("목표 마진율", cfg["targetMarginRate"], "이 이상이면 판매"),
            ("상품당 포장비 (원)", cfg["packingCost_KRW"], "예시값")]
    srows = [[("역계산 설정 (노란 칸 수정 → 수집결과 자동 재계산)", "title")], []]
    for k, v, note in conf:
        srows.append([(k, "plain"), (v, "input"), (note, "note")])
    srows += [[], [], [("국제배송비 표 (무게 하한 g → 배송비 원, 예시값)", "plain")]]
    for lo, cost in [(0, 5500), (500, 7500), (1000, 10500), (2000, 15000), (3000, 19500)]:
        srows.append([(lo, "input_money"), (cost, "input_money")])
    srows += [[], [(f"수집일: {datetime.date.today()} / 키워드: {', '.join(cfg['keywords'])}", "note")]]

    out = [[("라쿠텐 일본 수요 (리뷰 많은 순)", "title")],
           [("노란 칸만 채우면 판정이 나와요. 네이버/다나와/쿠팡을 누르면 한국 검색어로 가격 비교가 열려요.", "note")],
           [], [(h, "head") for h in COLS]]
    for i, v in enumerate(rows):
        r = 5 + i
        Q = f'{X["한국 검색어 (수정 가능)"]}{r}'
        g = lambda n: f'{X[n]}{r}'
        row = {
            "순위": (i + 1, "plain"), "키워드": (", ".join(v["keywords"]), "plain"), "상품명 (일본어)": (v["name"], "plain"),
            "한국어 상품명": (v.get("name_ko", ""), "plain"), "일본 판매가 (엔)": (v["price"], "money"),
            "리뷰 수": (v["reviews"], "money"), "평점": (v["rating"], "plain"), "샵": (v["shop"], "plain"),
            "배송비 포함": (v["postage"], "plain"),
            "라쿠텐": (f'=HYPERLINK("{v["url"]}","열기")', "link"),
            "한국 검색어 (수정 가능)": (make_query(v), "input"),
            "한국 구매가 (원)": ("", "input_money"), "무게 (g)": ("", "input_money"), "일본 성분 확인": ("", "input"),
            "판매가 (원)": (f'={g("일본 판매가 (엔)")}*설정!$B$3', "money"),
            "수수료 (원)": (f'={g("판매가 (원)")}*설정!$B$4', "money"),
            "배송비 (원)": (f'=IF({g("무게 (g)")}="","",INDEX(설정!$B$11:$B$15,MATCH({g("무게 (g)")},설정!$A$11:$A$15,1)))', "money"),
            "포장+환율여유 (원)": (f'=설정!$B$7+{g("판매가 (원)")}*설정!$B$5', "money"),
            "순수익 (원)": (f'=IF(OR({g("한국 구매가 (원)")}="",{g("무게 (g)")}=""),"",{g("판매가 (원)")}-{g("한국 구매가 (원)")}-{g("수수료 (원)")}-{g("배송비 (원)")}-{g("포장+환율여유 (원)")})', "money"),
            "마진율": (f'=IF({g("순수익 (원)")}="","",{g("순수익 (원)")}/{g("판매가 (원)")})', "pct"),
            "판정": (f'=IF({g("순수익 (원)")}="","입력 필요",IF({g("일본 성분 확인")}<>"O","보류",IF({g("마진율")}>=설정!$B$6,"판매","패스")))', "plain"),
            "최대 구매가 (원)": (f'=IF({g("무게 (g)")}="","",{g("판매가 (원)")}*(1-설정!$B$4-설정!$B$5-설정!$B$6)-{g("배송비 (원)")}-설정!$B$7)', "money"),
        }
        for name, base, label in SEARCH:
            row[name] = (f'=IF({Q}="","",HYPERLINK("{base}"&_xlfn.ENCODEURL({Q}),"{label}"))', "link")
        out.append([row[n] for n in COLS])
    widths = {C[n]: w for n, w in {"순위": 6, "키워드": 14, "상품명 (일본어)": 36, "한국어 상품명": 36, "샵": 16,
                                   "한국 검색어 (수정 가능)": 28, "네이버쇼핑": 8, "다나와": 8, "쿠팡": 8, "라쿠텐": 7}.items()}
    for n, i in C.items():
        widths.setdefault(i, 11)
    last = 4 + max(len(rows), 1)
    XL.save(path, [{"name": "수집결과", "rows": out, "widths": widths, "freeze": (4, 4), "filter": f"A4:{XL.col_letter(len(COLS))}{last}"},
                   {"name": "설정", "rows": srows, "widths": {1: 22, 2: 12, 3: 40}}])


def write_csv(rows, path):
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as fp:
        wr = csv.writer(fp)
        wr.writerow(["순위", "키워드", "상품명", "한국어 상품명", "일본 판매가(엔)", "리뷰 수", "평점", "샵", "배송비 포함", "라쿠텐", "한국 검색어", "네이버쇼핑", "다나와", "쿠팡"])
        for i, v in enumerate(rows, 1):
            wr.writerow([i, ", ".join(v["keywords"]), v["name"], v.get("name_ko", ""), v["price"], v["reviews"], v["rating"], v["shop"], v["postage"], v["url"], make_query(v)] + list(search_urls(make_query(v)).values()))


def main():
    cfg = load_config()
    rows = collect(cfg)
    add_korean_names(cfg, rows)
    print(f"\n총 {len(rows)}개 상품 (리뷰 {cfg.get('minReviewCount', 0)}개 이상)")
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    sys.path.insert(0, BASE)
    try:
        import openpyxl  # noqa
        out = os.path.join(BASE, f"일본수요_{stamp}.xlsx")
        write_xlsx(cfg, rows, out)
    except ImportError:
        out = os.path.join(BASE, f"일본수요_{stamp}.xlsx")
        try:
            write_xlsx_lite(cfg, rows, out)
            print("(openpyxl 없이 기본 엑셀로 저장 — 링크·계산은 되고 판정 색상만 없어요)")
        except Exception as e:
            out = os.path.join(BASE, f"일본수요_{stamp}.csv")
            write_csv(rows, out)
            print(f"엑셀 저장 실패({e}) → CSV로 저장했어요.")
    print("저장:", out)


if __name__ == "__main__":
    main()
